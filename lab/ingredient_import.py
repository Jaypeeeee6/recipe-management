"""Excel import helpers for ingredients."""

from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO

from django.utils import timezone
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .models import (
    Category,
    Ingredient,
    IngredientPriceHistory,
    Supplier,
    TrialProductStatus,
)
from .utils import next_ingredient_code

# Matches IngredientForm.jsx dropdowns
STORAGE_OPTIONS = ["Freezer", "Chiller", "Room Temperature"]
UNIT_OPTIONS = ["kg", "g", "L", "ml", "pcs"]
STATUS_OPTIONS = ["trial", "approved"]
SECRET_OPTIONS = ["yes", "no"]

# Column headers for the template / import (row 1).
INGREDIENT_IMPORT_COLUMNS = [
    ("name", "Name *"),
    ("category", "Category *"),
    ("supplier", "Supplier"),
    ("batch_number", "Batch Number"),
    ("expiry_date", "Expiry Date (YYYY-MM-DD)"),
    ("received_date", "Received Date (YYYY-MM-DD)"),
    ("unit", "Unit"),
    ("quantity", "Quantity"),
    ("par_level", "Par Level"),
    ("price_per_unit", "Price Per Unit (OMR)"),
    ("notes", "Notes"),
    ("storage_conditions", "Storage Conditions"),
    ("shelf_life_days", "Shelf Life Days"),
    ("status", "Status (trial/approved)"),
    ("is_secret", "Secret (yes/no)"),
]

SAMPLE_ROW = {
    "name": "Example Flour",
    "category": "Flour",
    "supplier": "",
    "batch_number": "BATCH-001",
    "expiry_date": "2026-12-31",
    "received_date": date.today().isoformat(),
    "unit": "kg",
    "quantity": "25",
    "par_level": "10",
    "price_per_unit": "0.850",
    "notes": "Sample row — delete before importing real data",
    "storage_conditions": "Room Temperature",
    "shelf_life_days": "",
    "status": "trial",
    "is_secret": "no",
}

# How many data rows get dropdowns (row 2 = sample, then blank rows for entry).
IMPORT_DROPDOWN_ROWS = 500


def _col_letter(key: str) -> str:
    for idx, (col_key, _) in enumerate(INGREDIENT_IMPORT_COLUMNS, start=1):
        if col_key == key:
            return get_column_letter(idx)
    raise KeyError(key)


def _write_list_column(ws, col_idx: int, header: str, values: list[str]) -> str | None:
    """Write a named list column on the Lists sheet. Returns Excel range like $A$2:$A$10."""
    ws.cell(row=1, column=col_idx, value=header)
    cleaned = [str(v).strip() for v in values if v and str(v).strip()]
    if not cleaned:
        return None
    for row_idx, value in enumerate(cleaned, start=2):
        ws.cell(row=row_idx, column=col_idx, value=value)
    letter = get_column_letter(col_idx)
    return f"Lists!${letter}$2:${letter}${len(cleaned) + 1}"


def _add_list_validation(ws, formula: str, col_letter: str, allow_blank=True):
    if not formula:
        return
    dv = DataValidation(
        type="list",
        formula1=formula,
        allow_blank=allow_blank,
        showDropDown=False,  # False = show the dropdown arrow in Excel
        showErrorMessage=True,
        errorTitle="Invalid value",
        error="Please choose a value from the dropdown list.",
        promptTitle="Choose",
        prompt="Select a value from the list",
    )
    dv.add(f"{col_letter}2:{col_letter}{IMPORT_DROPDOWN_ROWS}")
    ws.add_data_validation(dv)


def build_ingredient_import_template(
    *,
    categories: list[str] | None = None,
    suppliers: list[str] | None = None,
) -> bytes:
    category_names = list(categories or [])
    supplier_names = list(suppliers or [])
    if not category_names:
        category_names = list(
            Category.objects.order_by("sort_order", "name").values_list("name", flat=True)
        )
    if not supplier_names:
        supplier_names = list(
            Supplier.objects.order_by("company_name").values_list("company_name", flat=True)
        )

    sample = dict(SAMPLE_ROW)
    if category_names and sample["category"] not in category_names:
        sample["category"] = category_names[0]
    if supplier_names:
        sample["supplier"] = supplier_names[0]

    wb = Workbook()
    ws = wb.active
    ws.title = "Ingredients"
    headers = [label for _, label in INGREDIENT_IMPORT_COLUMNS]
    ws.append(headers)
    ws.append([sample.get(key, "") for key, _ in INGREDIENT_IMPORT_COLUMNS])
    for col in ws.columns:
        letter = col[0].column_letter
        ws.column_dimensions[letter].width = 22

    lists = wb.create_sheet("Lists")
    ranges = {
        "category": _write_list_column(lists, 1, "Categories", category_names),
        "supplier": _write_list_column(lists, 2, "Suppliers", supplier_names),
        "unit": _write_list_column(lists, 3, "Units", UNIT_OPTIONS),
        "storage_conditions": _write_list_column(lists, 4, "Storage", STORAGE_OPTIONS),
        "status": _write_list_column(lists, 5, "Status", STATUS_OPTIONS),
        "is_secret": _write_list_column(lists, 6, "Secret", SECRET_OPTIONS),
    }
    for col_idx in range(1, 7):
        lists.column_dimensions[get_column_letter(col_idx)].width = 22

    for key, formula in ranges.items():
        _add_list_validation(ws, formula, _col_letter(key), allow_blank=(key != "category"))

    # Hide Lists so users only see Ingredients + Instructions; dropdowns still work.
    lists.sheet_state = "hidden"

    info = wb.create_sheet("Instructions")
    for line in (
        "Keep the header row on the Ingredients sheet exactly as shown.",
        "Use the dropdowns for Category, Supplier, Unit, Storage Conditions, Status, and Secret.",
        "Required: Name * and Category * (pick a category from the dropdown).",
        "Supplier is optional — leave blank or pick from the dropdown.",
        "Dates: YYYY-MM-DD (example 2026-12-31).",
        "Status: trial (default) or approved.",
        "Secret: yes/no (default no). Only Admin/IT can import secret=yes rows.",
        "Storage: Freezer, Chiller, or Room Temperature.",
        "Ingredient codes are assigned automatically — do not add a Code column.",
        "Delete the sample row before importing real data.",
        "Re-download this template after adding categories or suppliers so dropdowns stay current.",
    ):
        info.append([line])
    info.column_dimensions["A"].width = 90

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _cell(row, idx):
    if idx >= len(row):
        return ""
    value = row[idx]
    if value is None:
        return ""
    return value


def _as_str(value):
    if value is None:
        return ""
    if isinstance(value, float) and value == int(value):
        return str(int(value))
    return str(value).strip()


def _as_bool(value, default=False):
    text = _as_str(value).lower()
    if not text:
        return default
    if text in ("1", "true", "yes", "y", "secret"):
        return True
    if text in ("0", "false", "no", "n"):
        return False
    return default


def _as_decimal(value, default=Decimal("0")):
    text = _as_str(value).replace(",", "")
    if not text:
        return default
    try:
        return Decimal(text)
    except (InvalidOperation, ValueError):
        raise ValueError(f"Invalid number: {value!r}")


def _as_int(value):
    text = _as_str(value)
    if not text:
        return None
    try:
        return int(float(text))
    except (TypeError, ValueError):
        raise ValueError(f"Invalid integer: {value!r}")


def _as_date(value):
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _as_str(value)
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"Invalid date (use YYYY-MM-DD): {value!r}")


def _normalize_headers(header_row):
    mapping = {}
    aliases = {
        "name": ("name", "name *", "ingredient", "ingredient name"),
        "category": ("category", "category *", "category name"),
        "supplier": ("supplier", "supplier name", "company"),
        "batch_number": ("batch number", "batch", "batch_number", "lot"),
        "expiry_date": ("expiry date (yyyy-mm-dd)", "expiry date", "expiry", "expiry_date"),
        "received_date": ("received date (yyyy-mm-dd)", "received date", "received", "received_date"),
        "unit": ("unit",),
        "quantity": ("quantity", "qty", "stock"),
        "par_level": ("par level", "par", "par_level"),
        "price_per_unit": ("price per unit (omr)", "price per unit", "price", "unit price", "price_per_unit"),
        "notes": ("notes", "note"),
        "storage_conditions": ("storage conditions", "storage", "storage_conditions"),
        "shelf_life_days": ("shelf life days", "shelf life", "shelf_life_days"),
        "status": ("status (trial/approved)", "status"),
        "is_secret": ("secret (yes/no)", "secret", "is_secret", "is secret"),
    }
    for idx, raw in enumerate(header_row):
        label = _as_str(raw).lower()
        if not label:
            continue
        for key, names in aliases.items():
            if label in names and key not in mapping:
                mapping[key] = idx
                break
    return mapping


def import_ingredients_from_excel(file_obj, *, allow_secret=False):
    """
    Import ingredients from an .xlsx file.
    Returns dict: created, skipped, errors (list of {row, message}).
    """
    wb = load_workbook(file_obj, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return {"created": 0, "skipped": 0, "errors": [{"row": 1, "message": "File is empty."}]}

    header_map = _normalize_headers(rows[0])
    if "name" not in header_map or "category" not in header_map:
        return {
            "created": 0,
            "skipped": 0,
            "errors": [
                {
                    "row": 1,
                    "message": "Missing required columns: Name * and Category *.",
                }
            ],
        }

    categories = {c.name.strip().lower(): c for c in Category.objects.all()}
    suppliers = {s.company_name.strip().lower(): s for s in Supplier.objects.all()}

    created = 0
    skipped = 0
    errors = []

    for row_number, row in enumerate(rows[1:], start=2):
        if not row or all(v is None or _as_str(v) == "" for v in row):
            skipped += 1
            continue

        try:
            name = _as_str(_cell(row, header_map["name"]))
            category_name = _as_str(_cell(row, header_map["category"]))
            if not name:
                raise ValueError("Name is required.")
            if not category_name:
                raise ValueError("Category is required.")
            notes_val = (
                _as_str(_cell(row, header_map["notes"])) if "notes" in header_map else ""
            )
            # Skip the template sample row if left in place
            if name.lower().startswith("example ") and "sample row" in notes_val.lower():
                skipped += 1
                continue

            category = categories.get(category_name.lower())
            if not category:
                raise ValueError(f"Unknown category: {category_name!r}")

            supplier = None
            if "supplier" in header_map:
                supplier_name = _as_str(_cell(row, header_map["supplier"]))
                if supplier_name:
                    supplier = suppliers.get(supplier_name.lower())
                    if not supplier:
                        raise ValueError(f"Unknown supplier: {supplier_name!r}")

            status_raw = _as_str(_cell(row, header_map["status"])) if "status" in header_map else "trial"
            status_key = status_raw.lower() or "trial"
            if status_key not in ("trial", "approved", "testing"):
                raise ValueError("Status must be trial or approved.")
            is_trial = status_key in ("trial", "testing")
            trial_status = (
                TrialProductStatus.TESTING if is_trial else TrialProductStatus.APPROVED
            )

            is_secret = False
            if "is_secret" in header_map:
                is_secret = _as_bool(_cell(row, header_map["is_secret"]), False)
            if is_secret and not allow_secret:
                raise ValueError("Only Admin/IT can import secret ingredients.")

            def col(key, default=""):
                if key not in header_map:
                    return default
                return _cell(row, header_map[key])

            quantity = _as_decimal(col("quantity"), Decimal("0"))
            par_level = _as_decimal(col("par_level"), Decimal("0"))
            price = _as_decimal(col("price_per_unit"), Decimal("0"))
            unit = _as_str(col("unit")) or "kg"
            code = next_ingredient_code(category)

            ingredient = Ingredient.objects.create(
                code=code,
                name=name,
                category=category,
                supplier=supplier,
                batch_number=_as_str(col("batch_number")),
                expiry_date=_as_date(col("expiry_date")),
                received_date=_as_date(col("received_date")),
                unit=unit,
                quantity=quantity,
                par_level=par_level,
                price_per_unit=price,
                notes=_as_str(col("notes")),
                storage_conditions=_as_str(col("storage_conditions")),
                shelf_life_days=_as_int(col("shelf_life_days")),
                is_secret=is_secret,
                is_trial=is_trial,
                trial_status=trial_status,
            )
            if price:
                IngredientPriceHistory.objects.create(
                    ingredient=ingredient,
                    price=price,
                    recorded_at=timezone.localdate(),
                )
            created += 1
        except Exception as exc:  # noqa: BLE001 — collect per-row errors
            errors.append({"row": row_number, "message": str(exc)})

    return {"created": created, "skipped": skipped, "errors": errors}
